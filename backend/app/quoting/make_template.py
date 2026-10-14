"""Generează templates/oferta.docx (șablon docxtpl). Rulează o dată:
    python -m app.quoting.make_template
Apoi șablonul poate fi editat în Word (logo, fonturi, antet) — păstrează etichetele {{ }} / {% %}.
Documentul e pentru CLIENT: nu conține prețuri de achiziție, adaos, furnizori sau curs."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

TEMPLATE = Path(__file__).parent / "templates" / "oferta.docx"
COLUMNS = ["Nr.", "Cod", "Descriere", "Producător", "Cant.", "Preț unitar (EUR)", "Valoare (EUR)"]
FIELDS = ["l.nr", "l.code", "l.description", "l.manufacturer", "l.quantity", "l.unit_price", "l.total"]


def build(path=TEMPLATE):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_paragraph("[Denumirea firmei] · [CUI] · [Adresă] · [Contact]")
    doc.add_heading("OFERTĂ COMERCIALĂ nr. {{ offer_number }}", level=1)
    doc.add_paragraph("Către: {{ client_name }}")
    doc.add_paragraph("Data: {{ offer_date }}    Valabilitate: {{ validity_days }} zile (până la {{ valid_until }})")

    table = doc.add_table(rows=4, cols=len(COLUMNS))
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, COLUMNS):
        cell.text = title
    table.rows[1].cells[0].text = "{%tr for l in lines %}"
    for cell, f in zip(table.rows[2].cells, FIELDS):
        cell.text = "{{ " + f + " }}"
    table.rows[3].cells[0].text = "{%tr endfor %}"

    doc.add_paragraph("Total ofertă: {{ total }} EUR (fără TVA)").runs[0].bold = True

    doc.add_heading("Condiții comerciale", level=2)
    for line in [
        "Prețurile sunt ferme și exprimate exclusiv în EUR, fără TVA.",
        "Condiții de livrare: {{ delivery_terms }}.",
        "Penalități de întârziere: {{ penalty_per_week }}% din valoarea comenzii pentru fiecare săptămână "
        "de întârziere, în limita a maximum {{ penalty_cap }}% din valoarea comenzii.",
        "Disponibilitatea pieselor de schimb: garantată minimum {{ spare_parts_years }} de ani de la livrare.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Conformitate", level=2)
    doc.add_paragraph("{% for c in certificates %}Certificare {{ c.standard }}, nr. {{ c.number }}, "
                      "valabilă până la {{ c.valid_until }}.\n{% endfor %}")
    doc.add_paragraph("{% if afer %}Agremente tehnice AFER pentru produsele din import: "
                      "{% for a in afer %}{{ a.code }} – {{ a.agreement }}{% if not loop.last %}; {% endif %}"
                      "{% endfor %}.{% endif %}")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build())
