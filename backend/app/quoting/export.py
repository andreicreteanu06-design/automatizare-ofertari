"""Export: Excel = calcul INTERN (costuri, curs, adaos); Word = draft de ofertă pentru CLIENT."""
import io
from datetime import timedelta
from decimal import Decimal

from docxtpl import DocxTemplate
from openpyxl import Workbook
from openpyxl.styles import Font

from app.quoting.make_template import TEMPLATE

MONEY, RATE = "#,##0.00", "#,##0.0000"


def fmt_eur(value):
    """1234.5 -> '1.234,50' (format românesc)."""
    return f"{value:,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def _fmt_qty(q):
    return f"{q.normalize():f}".replace(".", ",")


def excel(offer, calc, failures=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Calcul intern"
    ws.append([f"CALCUL INTERN — CONFIDENȚIAL — oferta {offer.offer_number} — {offer.client_name}"])
    ws["A1"].font = Font(bold=True, color="C00000")
    ws.append([f"Curs BNR din {calc.fx.date.isoformat()}: 1 EUR = {calc.fx.rate('EUR')} RON; "
               f"adaos {calc.markup_percent}% peste costul total livrat"])
    ws.append([])
    header = ["Cod", "Descriere", "Furnizor", "Producător", "Cant.", "Preț achiziție", "Monedă",
              "Preț EUR/buc", "Transport EUR/buc", "Taxă vamală %", "Taxă vamală EUR/buc",
              "Cost livrat EUR/buc", "Preț vânzare EUR/buc", "Total vânzare EUR", "Marjă EUR"]
    ws.append(header)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for c in calc.lines:
        l = c.line
        ws.append([l.code, l.description, l.supplier, l.manufacturer, l.quantity, l.unit_price,
                   l.price_currency.upper(), c.price_eur, c.transport_eur, l.duty_rate_percent, c.duty_eur,
                   c.landed_eur, c.sale_unit_eur, c.sale_total_eur, c.margin_eur])
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", "", "", "", "", "", calc.total_landed_eur, "",
               calc.total_sale_eur, calc.total_margin_eur])
    ws.append(["Marjă raportată la vânzare %", "", "", "", "", "", "", "", "", "", "", "", "",
               calc.margin_on_sale_percent])
    for row in ws.iter_rows(min_row=5, min_col=6, max_col=15):
        for cell in row:
            if isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = RATE if cell.column in (8, 9, 11, 12) else MONEY
    for col, width in zip("ABCDEFGHIJKLMNO", [16, 36, 24, 18, 8, 14, 8, 14, 16, 12, 18, 18, 18, 18, 14]):
        ws.column_dimensions[col].width = width

    fx = wb.create_sheet("Curs BNR")
    fx.append(["Monedă", "RON / 1 unitate", "Data"])
    used = sorted({"EUR", *(l.price_currency.upper() for l in offer.lines),
                   *(l.transport_currency.upper() for l in offer.lines)} - {"RON"})
    for cur in used:
        fx.append([cur, calc.fx.rate(cur), calc.fx.date.isoformat()])

    chk = wb.create_sheet("Checklist")
    chk.append(["Cerință", "Status"])
    chk.append(["Checklist conformitate", "OK" if not failures else "BLOCAT"])
    for f in failures:
        chk.append([f.code, f.message])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def word(offer, calc):
    tpl = DocxTemplate(TEMPLATE)
    tpl.render({
        "offer_number": offer.offer_number,
        "client_name": offer.client_name,
        "offer_date": offer.offer_date.strftime("%d.%m.%Y"),
        "validity_days": offer.validity_days,
        "valid_until": (offer.offer_date + timedelta(days=offer.validity_days)).strftime("%d.%m.%Y"),
        "lines": [{"nr": i, "code": c.line.code, "description": c.line.description,
                   "manufacturer": c.line.manufacturer or "-", "quantity": _fmt_qty(c.line.quantity),
                   "unit_price": fmt_eur(c.sale_unit_eur), "total": fmt_eur(c.sale_total_eur)}
                  for i, c in enumerate(calc.lines, 1)],
        "total": fmt_eur(calc.total_sale_eur),
        "delivery_terms": offer.delivery_terms,
        "penalty_per_week": _fmt_qty(offer.penalty_percent_per_week),
        "penalty_cap": _fmt_qty(offer.penalty_cap_percent),
        "spare_parts_years": offer.spare_parts_years,
        "certificates": [{"standard": c.standard, "number": c.number, "valid_until": c.valid_until.strftime("%d.%m.%Y")}
                         for c in offer.certificates],
        "afer": [{"code": l.code, "agreement": l.afer_agreement} for l in offer.lines if l.imported],
    })
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()
