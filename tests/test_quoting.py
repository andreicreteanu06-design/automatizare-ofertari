import io
from datetime import date
from decimal import Decimal

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app, get_fx
from app.quoting import compliance
from app.quoting.engine import FxError, FxSnapshot, calc_line, calc_offer
from app.quoting.fx import bnr_snapshot
from app.quoting.models import Offer, OfferLine

D = Decimal
FX = FxSnapshot(date(2026, 10, 13), {"EUR": D("5.0000"), "USD": D("4.5000"), "HUF": D("0.0125")})

LINE_A = {"code": "3EST000123", "description": "Suport boghiu", "supplier": "Feroparts SRL",
          "manufacturer": "Knorr-Bremse", "quantity": "4", "unit_price": "555", "price_currency": "EUR",
          "transport_cost": "40", "transport_currency": "EUR"}
LINE_B = {"code": "DIN933-M12X40", "description": "Surub M12x40", "supplier": "Bolt Inc", "quantity": "10",
          "unit_price": "100", "price_currency": "usd", "transport_cost": "500", "transport_currency": "RON",
          "duty_rate_percent": "2.7", "imported": True, "afer_agreement": "AT-0123/2025"}
VALID = {
    "offer_number": "OF-2026/045", "client_name": "Client Feroviar SA", "offer_date": "2026-10-13",
    "markup_percent": "20", "currency": "EUR", "firm_prices": True, "spare_parts_years": 30,
    "spare_parts_reference": "Scrisoare Feroparts 12/2026", "include_penalty_clause": True,
    "certificates": [
        {"supplier": "Feroparts SRL", "standard": "ISO/TS 22163", "number": "IRIS-123", "valid_until": "2027-05-01"},
        {"supplier": "BOLT  INC", "standard": "ISO 9001", "number": "Q-9", "valid_until": "2026-10-13"}],
    "lines": [LINE_A, LINE_B],
}


def offer(**changes):
    return Offer(**{**VALID, **changes})


# --- Calcul (verificat de mână) ------------------------------------------------------

def test_landed_cost_eur_line():
    c = calc_line(OfferLine(**LINE_A), FX, D(20))
    # 555 + 40/4 transport, fără taxă vamală -> 565; +20% -> 678.00; × 4
    assert (c.price_eur, c.transport_eur, c.duty_eur, c.landed_eur) == (D(555), D(10), D(0), D(565))
    assert (c.sale_unit_eur, c.sale_total_eur, c.margin_eur) == (D("678.00"), D("2712.00"), D("452.00"))


def test_landed_cost_import_line_usd_ron_duty():
    c = calc_line(OfferLine(**LINE_B), FX, D(20))
    # 100 USD × 4.5 / 5 = 90 EUR; 500 RON / 5 = 100 EUR / 10 buc = 10; taxă (90+10) × 2.7% = 2.7
    assert (c.price_eur, c.transport_eur, c.duty_eur, c.landed_eur) == (D(90), D(10), D("2.7"), D("102.7"))
    assert (c.sale_unit_eur, c.sale_total_eur, c.margin_eur) == (D("123.24"), D("1232.40"), D("205.40"))


def test_offer_totals_and_real_margin():
    calc = calc_offer(offer(), FX)
    assert (calc.total_sale_eur, calc.total_landed_eur, calc.total_margin_eur) == (
        D("3944.40"), D("3287.00"), D("657.40"))
    assert calc.margin_on_sale_percent == D("16.67")  # 20% adaos ≠ 20% marjă


def test_rounding_half_up_once():
    c = calc_line(OfferLine(**{**LINE_A, "unit_price": "0.125", "transport_cost": "0"}), FX, D(0))
    assert c.sale_unit_eur == D("0.13")  # nu 0.12 (rotunjire bancară)


def test_huf_multiplier_and_ron():
    assert FX.to_eur(D(10000), "HUF") == D(25)  # 10000 × 0.0125 / 5
    assert FX.to_eur(D(50), "RON") == D(10)


def test_missing_currency_rate():
    with pytest.raises(FxError, match="GBP"):
        calc_line(OfferLine(**{**LINE_A, "price_currency": "GBP"}), FX, D(0))


# --- Checklist de conformitate -------------------------------------------------------

def codes(o):
    return [f.code for f in compliance.check(o)]


def test_valid_offer_passes():
    assert codes(offer()) == []


@pytest.mark.parametrize("changes,expected", [
    ({"lines": [LINE_A, {**LINE_B, "afer_agreement": None}]}, ["AFER_MISSING"]),
    ({"lines": [LINE_A, {**LINE_B, "afer_agreement": "  "}]}, ["AFER_MISSING"]),
    ({"certificates": VALID["certificates"][:1]}, ["CERT_MISSING"]),
    ({"certificates": [VALID["certificates"][0], {**VALID["certificates"][1], "valid_until": "2026-10-12"}]},
     ["CERT_EXPIRED"]),
    ({"spare_parts_years": 25}, ["SPARE_PARTS_30Y"]),
    ({"spare_parts_reference": ""}, ["SPARE_PARTS_30Y"]),
    ({"include_penalty_clause": False}, ["PENALTY_CLAUSE"]),
    ({"penalty_percent_per_week": "3"}, ["PENALTY_CLAUSE"]),
    ({"penalty_cap_percent": "20"}, ["PENALTY_CLAUSE"]),
    ({"currency": "USD"}, ["CURRENCY_NOT_EUR"]),
    ({"firm_prices": False}, ["PRICES_NOT_FIRM"]),
])
def test_each_requirement_blocks(changes, expected):
    assert codes(offer(**changes)) == expected


def test_domestic_line_needs_no_afer():
    assert codes(offer(lines=[LINE_A])) == []


def test_unchecked_offer_lists_everything():
    bare = offer(firm_prices=False, spare_parts_years=0, include_penalty_clause=False, certificates=[],
                 lines=[LINE_A, {**LINE_B, "afer_agreement": None}])
    assert codes(bare) == ["AFER_MISSING", "CERT_MISSING", "CERT_MISSING", "SPARE_PARTS_30Y",
                           "PENALTY_CLAUSE", "PRICES_NOT_FIRM"]


# --- API: calcul + export blocat de checklist ---------------------------------------

@pytest.fixture
def client():
    calls = []

    async def fake_fx(on):
        calls.append(on)
        return FX
    app.dependency_overrides[get_fx] = lambda: fake_fx
    yield TestClient(app), calls
    app.dependency_overrides.clear()


def test_calculate_endpoint(client):
    c, calls = client
    body = c.post("/offers/calculate", json={**VALID, "firm_prices": False}).json()
    assert calls == [date(2026, 10, 13)]
    assert body["total_sale_eur"] == "3944.40" and body["fx"] == {"source": "BNR", "date": "2026-10-13", "eur_ron": "5.0000"}
    assert body["lines"][1]["landed_eur"] == "102.7000"
    assert body["compliance"] == {"ok": False, "failures": [
        {"code": "PRICES_NOT_FIRM", "message": "prețurile trebuie confirmate ca ferme"}]}


def test_fx_date_override(client):
    c, calls = client
    c.post("/offers/calculate", json={**VALID, "fx_date": "2026-10-01"})
    assert calls == [date(2026, 10, 1)]


@pytest.mark.parametrize("kind", ["xlsx", "docx"])
def test_export_refused_when_checklist_fails(client, kind):
    c, calls = client
    r = c.post(f"/offers/export/{kind}", json={**VALID, "include_penalty_clause": False, "currency": "RON"})
    assert r.status_code == 409
    assert [f["code"] for f in r.json()["detail"]["failures"]] == ["PENALTY_CLAUSE", "CURRENCY_NOT_EUR"]
    assert calls == []  # refuzat înainte de curs și de generare


def test_export_xlsx_internal(client):
    c, _ = client
    r = c.post("/offers/export/xlsx", json=VALID)
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="calcul_intern_OF-2026_045.xlsx"'
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Calcul intern"]
    assert "CONFIDENȚIAL" in ws["A1"].value
    rows = {row[0]: row for row in ws.iter_rows(min_row=5, values_only=True) if row[0]}
    assert rows["DIN933-M12X40"][11] == pytest.approx(102.7) and rows["DIN933-M12X40"][12] == pytest.approx(123.24)
    assert rows["TOTAL"][13] == pytest.approx(3944.40)
    assert [r for r in wb["Curs BNR"].iter_rows(values_only=True)][1:] == [
        ("EUR", pytest.approx(5.0), "2026-10-13"), ("USD", pytest.approx(4.5), "2026-10-13")]
    assert wb["Checklist"]["B2"].value == "OK"


def docx_text(content):
    doc = Document(io.BytesIO(content))
    cells = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_export_docx_for_client(client):
    c, _ = client
    r = c.post("/offers/export/docx", json=VALID)
    assert r.status_code == 200
    text = docx_text(r.content)
    for expected in ["OFERTĂ COMERCIALĂ nr. OF-2026/045", "Către: Client Feroviar SA", "până la 12.11.2026",
                     "678,00", "2.712,00", "1.232,40", "Total ofertă: 3.944,40 EUR",
                     "ferme și exprimate exclusiv în EUR",
                     "5% din valoarea comenzii pentru fiecare săptămână", "maximum 30%",
                     "minimum 30 de ani", "ISO/TS 22163, nr. IRIS-123", "DIN933-M12X40 – AT-0123/2025"]:
        assert expected in text, expected


def test_docx_has_no_internal_data(client):
    c, _ = client
    text = docx_text(c.post("/offers/export/docx", json=VALID).content)
    for internal in ["555", "565", "102,7", "Feroparts", "Bolt Inc", "adaos", "Marjă", "marj", "RON", "5,0000",
                     "{{", "{%"]:
        assert internal not in text, internal


def test_fx_unavailable_is_502(client):
    async def down(on):
        raise httpx.ConnectError("gateway down")
    app.dependency_overrides[get_fx] = lambda: down
    c, _ = client
    assert c.post("/offers/export/docx", json=VALID).status_code == 502


@pytest.mark.parametrize("bad", [{"lines": []}, {"markup_percent": "-1"},
                                 {"lines": [{**LINE_A, "quantity": "0"}]},
                                 {"lines": [{**LINE_A, "price_currency": "EURO"}]}])
def test_offer_validation(client, bad):
    c, _ = client
    assert c.post("/offers/calculate", json={**VALID, **bad}).status_code == 422


# --- Clientul FX al backend-ului: vorbește doar cu research-gateway ------------------

def test_bnr_snapshot_via_gateway(monkeypatch):
    monkeypatch.setenv("RESEARCH_GATEWAY_URL", "http://research-gateway:8000")
    seen = []

    def handler(request):
        seen.append(str(request.url))
        return httpx.Response(200, json={"source": "BNR", "date": "2026-10-10",
                                         "rates": {"EUR": "4.9763", "HUF": "0.012500"}})
    import asyncio
    snap = asyncio.run(bnr_snapshot(date(2026, 10, 11), transport=httpx.MockTransport(handler)))
    assert seen == ["http://research-gateway:8000/fx/bnr?date=2026-10-11"]
    assert snap == FxSnapshot(date(2026, 10, 10), {"EUR": D("4.9763"), "HUF": D("0.012500")})


def test_bnr_snapshot_rejects_bad_payload(monkeypatch):
    monkeypatch.setenv("RESEARCH_GATEWAY_URL", "http://research-gateway:8000")
    import asyncio
    t = httpx.MockTransport(lambda r: httpx.Response(200, json={"date": "2026-10-10", "rates": {}}))
    with pytest.raises(FxError):
        asyncio.run(bnr_snapshot(date(2026, 10, 11), transport=t))
