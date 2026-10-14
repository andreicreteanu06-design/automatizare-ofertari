"""mail-classify: carantină -> data/inbox/{invoices,tech,pricelists,unsorted}. DOAR rețeaua internă.

Ordinea deciziei:
  1. determinist: .xml cu namespace UBL Invoice-2/CreditNote-2 -> invoices;
     .zip care conține un astfel de XML (arhiva SPV) -> invoices; alt .xml/.zip -> unsorted;
  2. PDF / Excel / CSV: Laya (decizie „choice”, probabilități calibrate) pe numele fișierului,
     subiectul e-mailului și începutul textului. Sub prag -> unsorted (revizuire manuală).
     O factură PDF -> unsorted: parserul e-Factura acceptă doar XML/ZIP.
Laya indisponibil => fișierul rămâne în carantină și e reîncercat (nu ghicim).
Conținutul atașamentelor e nesigur: Laya doar clasifică (nu execută nimic), iar o
clasificare greșită e prinsă de validarea parserelor de după.
"""
import asyncio
import io
import json
import logging
import os
import shutil
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import httpx

log = logging.getLogger("mail-classify")

FOLDERS = {"invoice": "invoices", "datasheet": "tech", "price_list": "pricelists"}
UNSORTED = "unsorted"
MIN_CONFIDENCE = float(os.environ.get("LAYA_MIN_CONFIDENCE", "0.7"))
STATE_CHARS = 4000
QUESTION = {"doc_type": {
    "type": "choice",
    "instructions": "What kind of business document is this e-mail attachment?",
    "criteria": {"invoice": "an invoice or credit note billing delivered goods",
                 "datasheet": "a technical datasheet, drawing, certificate, test report or product specification",
                 "price_list": "a supplier price list, quotation or commercial offer with prices",
                 "other": "anything else"}}}


def is_ubl(head):
    return b"Invoice-2" in head or b"CreditNote-2" in head


def zip_has_ubl(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                if info.filename.lower().endswith(".xml"):
                    with z.open(info) as f:  # citim doar începutul: fără risc de zip-bomb
                        if is_ubl(f.read(4096)):
                            return True
    except zipfile.BadZipFile:
        pass
    return False


def deterministic(name, data):
    """'invoice' / UNSORTED pentru .xml/.zip; None = are nevoie de Laya."""
    ext = Path(name).suffix.lower()
    if ext == ".xml":
        return "invoice" if is_ubl(data[:4096]) else UNSORTED
    if ext == ".zip":
        return "invoice" if zip_has_ubl(data) else UNSORTED
    return None


def content_text(name, data):
    """Începutul textului, best-effort: un fișier stricat/malițios dă text gol, nu excepție."""
    ext = Path(name).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages[:2])[:STATE_CHARS]
        if ext == ".xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
            rows = (" | ".join(str(c) for c in row if c is not None) for row in ws.iter_rows(max_row=25, values_only=True))
            return "\n".join(rows)[:STATE_CHARS]
        if ext == ".csv":
            return data[:STATE_CHARS].decode("utf-8", errors="replace")
    except Exception as e:  # noqa: BLE001 — orice eroare de parsare => clasificăm doar după nume/subiect
        log.info("text indisponibil pentru %s: %s", name, type(e).__name__)
    return ""


async def laya_classify(state, url=None, transport=None):
    async with httpx.AsyncClient(timeout=60, transport=transport) as client:
        r = await client.post(url or os.environ["LAYA_URL"],
                              json={"model": "laya", "state": state, "questions": QUESTION})
        r.raise_for_status()
        answer = r.json()["answers"]["doc_type"]
    label = answer["choice"]
    return label, float(answer["probabilities"][label])


def decide(label, confidence):
    if label in ("datasheet", "price_list") and confidence >= MIN_CONFIDENCE:
        return label
    return UNSORTED  # „other”, factură PDF, sau nesigur


async def process_once(quarantine, inbox, classify=laya_classify):
    """Procesează toate perechile (fișier, sidecar) complete. Întoarce [(fișier, folder)]."""
    moved = []
    for sidecar in sorted(quarantine.glob("*.json")):
        path = sidecar.with_suffix("")  # abc_x.pdf.json -> abc_x.pdf
        if not path.exists():
            continue
        meta = json.loads(sidecar.read_text("utf-8"))
        data = path.read_bytes()
        label, confidence = deterministic(path.name, data), 1.0
        if label is None:
            state = (f"File name: {meta.get('original_name', path.name)}\nE-mail subject: {meta.get('subject', '')}\n"
                     f"Content:\n{content_text(path.name, data)}")[:STATE_CHARS]
            try:
                laya_label, confidence = await classify(state)
            except (httpx.HTTPError, KeyError, ValueError) as e:
                log.warning("Laya indisponibil pentru %s (%s): rămâne în carantină", path.name, type(e).__name__)
                continue
            label = decide(laya_label, confidence)
        folder = FOLDERS.get(label, UNSORTED)
        dest = inbox / folder / path.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        # Carantina și inbox-ul sunt volume diferite: mutarea e copiere, nu rename atomic.
        # .part + replace => watcher-ul de ingestie nu vede niciodată un fișier pe jumătate scris.
        tmp = dest.with_name(dest.name + ".part")
        shutil.copyfile(path, tmp)
        os.replace(tmp, dest)  # numele începe cu hash-ul conținutului: aceeași cale = același fișier
        path.unlink()
        with open(inbox / "_mail_log.jsonl", "a", encoding="utf-8") as logf:
            logf.write(json.dumps({**meta, "file": f"{folder}/{path.name}", "label": label,
                                   "confidence": round(confidence, 4),
                                   "classified_at": datetime.now(timezone.utc).isoformat()}, ensure_ascii=False) + "\n")
        sidecar.unlink()
        moved.append((path.name, folder))
    return moved


def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(name)s %(message)s")
    quarantine = Path(os.environ.get("MAIL_QUARANTINE", "/quarantine"))
    inbox = Path(os.environ.get("MAIL_INBOX", "/data/inbox"))
    interval = int(os.environ.get("CLASSIFY_POLL_SECONDS", "30"))
    while True:
        for name, folder in asyncio.run(process_once(quarantine, inbox)):
            log.info("%s -> %s", name, folder)
        time.sleep(interval)


if __name__ == "__main__":
    main()
